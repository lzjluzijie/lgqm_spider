#!/usr/bin/env python
# -*- encoding: utf-8 -*-
'''
@File    :   lpr_spider.py
@Time    :   2019/08/27 20:07:19
@Author  :   Jeffrey Wang
@Version :   1.0
@Contact :   shwangjj@163.com
@Desc    :   LPR 数据爬虫实现代码
'''

import requests
import openpyxl
import xlrd
import datetime, time


URL_EXCEL= "http://www.shibor.org/shibor/web/html/downLoad.html?nameNew=Historical_LPR_Data_{year}_{month}.xls&nameOld=LPR%CA%FD%BE%DD2019_8.xls&shiborSrc=http%3A%2F%2Fwww.shibor.org%2Fshibor%2F&downLoadPath=data"


def spider_server_start():
    """
    下载服务器启动

    执行以下操作：
    1. 网上爬取最新的数据（最新判断依据是当前服务器日期）
    2. 同数据库检查，是否有新数据
    3. 如有新数据则更新数据库
    """
    # TODO 待实现
    pass


def load_history(year_month):
    """
    更新历史数据
    
    Arguments:
        year_month {string} -- yyyy或yyyyMM格式日期
    """
    # EFFECTS：
    # 0. 判断是否是2019.8之前
    # 1. 下载文件（如果是年，则为多个）
    # 2. 解析成pandas.dataframe（仅仅这个月或）
    # 3. 更新数据库

    # # 下载文件
    # url = URL_EXCEL.format( year=2019, month=8 )
    # r = requests.get( url )
    # with open("c:\\foo\\lpr.xls", mode='wb') as f:
    #     f.write(r.content)
    # # FIXME 下载失败处理！

    # 解析
    # https://www.cnblogs.com/insane-Mr-Li/p/9092619.html
    wb = xlrd.open_workbook( "c:\\foo\\lpr.xls" )
    ws = wb.sheets()[0] 
    # print( ws.name )
    # 遍历：
    for row in range(1, ws.nrows):
        # cell = ws.cell(row, 0)
        # print( ws.cell_type(row, 0) )
        if ws.cell_type(row, 0) == 3:
            dt_cell = xlrd.xldate_as_tuple(ws.cell(row, 0).value, wb.datemode)
            # date_tmp = datetime.datetime(dt_cell)
            date_tmp = datetime.date(*dt_cell[:3]).strftime("%Y%m%d")  
            print( date_tmp )
        # if cell.
        # dt_cell = xldate_as_tuple(ws.cell(row, 0), 0)
        # # print( ws.cell(row, 0).value )
        # print( dt_cell )


    # TODO 待实现
    pass


def get(date=None):
    """
    按日期查询LPR数据
    
    Keyword Arguments:
        date {string} -- yyyyMMdd格式日期，默认为最新日期 (default: {None})
    """
    # TODO 待实现
    pass



if __name__ == "__main__":
    load_history('201907')



